#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
脚本: scripts/record_2024q1_simulation_workbook.py
功能:
  1. 向 outputs/阶段A_框架初版/A07_预测日志.xlsx 写入 2024-01-19 与 2024-03-08 预测快照
  2. 向 outputs/阶段A_框架初版/A08_错误归因.xlsx 写入对应的行情兑现复盘与归因
  3. 向 outputs/阶段A_框架初版/A09_研究假设库.xlsx 更新 ARCH-H-007, ARCH-H-005, H-002, H-003 的验证记录
"""

import openpyxl
from datetime import datetime
from pathlib import Path

BASE_DIR = Path(__file__).resolve().parent.parent
A07_PATH = BASE_DIR / "outputs" / "阶段A_框架初版" / "A07_预测日志.xlsx"
A08_PATH = BASE_DIR / "outputs" / "阶段A_框架初版" / "A08_错误归因.xlsx"
A09_PATH = BASE_DIR / "outputs" / "阶段A_框架初版" / "A09_研究假设库.xlsx"

def update_a07():
    print("[+] 更新 A07_预测日志.xlsx ...")
    wb = openpyxl.load_workbook(A07_PATH)
    ws = wb['预测日志']

    # 预测记录 1: 2024-01-19
    pred_1 = [
        "PRED_20240119_01",
        datetime(2024, 1, 19, 18, 0, 0),
        "V0.2",
        "研究/案例/2024Q1_钢材成材走弱与原料负反馈踩踏推演报告.md",
        "Neutral_Weak",
        "Bearish",
        -0.45,
        "Weak",
        "Bullish_Speculation",
        "Weak_Contango",
        "Bearish_Hidden_Divergence",
        "Medium",
        "High",
        "Active",
        "冬储累库超季节性而表需冰冻；钢厂即期毛利深度亏损（-130元/吨），铁矿高位估值缺乏实物支撑，盘面强预期面临节后开工证伪风险",
        "原料高估值与成材弱现实隐性背离",
        "向下破位修正",
        15,
        35,
        "宏观节后出台超预期大规模刺激政策；节后下游开工复苏超预期",
        "螺纹05合约有效突破4050元/吨且现货基差大幅走强",
        "OBS_PROD_R01_20240119;OBS_INV_R02_20240119;FEAT_FBS_SCORE_20240119",
        "节后首周周度库存与表需发布(2024-02-18)",
        None
    ]

    # 预测记录 2: 2024-03-08
    pred_2 = [
        "PRED_20240308_01",
        datetime(2024, 3, 8, 18, 0, 0),
        "V0.2",
        "研究/案例/2024Q1_钢材成材走弱与原料负反馈踩踏推演报告.md",
        "Bearish",
        "Extreme_Bearish",
        -0.85,
        "Severe_Distress",
        "Short_Panic",
        "Super_Contango_Breakdown",
        "Bearish_Convergence",
        "High",
        "Extreme",
        "Active",
        "旺季不旺确立：螺纹周度表需仅200万吨（同比暴跌），厂库社库双爆仓；钢厂亏损超200元/吨倒逼大面积检修减产，压制铁水，铁矿港口累库破14100万吨，全产业链共振踩踏负反馈加速",
        "终端开工断崖+成本坍塌负反馈共振",
        "主跌浪向下破位",
        10,
        20,
        "政策紧急出台严厉粗钢行政平控压减限产提振成材利润；铁矿突发供应中断",
        "螺纹05合约反弹重回3750元/吨上方且铁水产量单周回升超5万吨",
        "FEAT_DAPP_RB_20240308;OBS_INV_R02_20240308;OBS_INV_M03_20240308;FEAT_FBS_SCORE_20240308",
        "螺纹跌破3500或钢厂自律减产扩围",
        None
    ]

    for col_idx, val in enumerate(pred_1, start=1):
        ws.cell(row=6, column=col_idx, value=val)
    for col_idx, val in enumerate(pred_2, start=1):
        ws.cell(row=7, column=col_idx, value=val)

    wb.save(A07_PATH)
    print("    -> A07 成功写入 2 条预测快照记录 (Row 6, Row 7)")

def update_a08():
    print("[+] 更新 A08_错误归因.xlsx ...")
    wb = openpyxl.load_workbook(A08_PATH)
    ws = wb['复盘记录']

    # 复盘 1 (针对 PRED_20240119_01)
    rev_1 = [
        "REV_20240329_01",
        "PRED_20240119_01",
        datetime(2024, 3, 29, 18, 0, 0),
        datetime(2024, 1, 19, 0, 0, 0),
        datetime(2024, 2, 23, 0, 0, 0),
        "数据/原始/01_期货行情.csv;数据/加工/加工指标.csv",
        "节前期货在3900~3970高位抗跌震荡，节后首周出现短期时滞，随后迅速暴跌至3790，逻辑完全兑现",
        "完全正确",
        "未交易(基准模拟空头盈利显著)",
        "无逻辑错误(存在预期时滞)",
        "季节性淡季流动性粘性",
        "节前市场博弈宏观强预期，现货停工无法提供微观证伪证据，节后开工数据落地后价格断崖式下行",
        "投机资金利用春节空窗期博弈政策利好",
        "在淡季增加隐性背离的持仓时滞衰减因子，避免过早在左侧加满杠杆",
        "ARCH-H-007;H-003",
        None,
        "V0.2",
        "已结项"
    ]

    # 复盘 2 (针对 PRED_20240308_01)
    rev_2 = [
        "REV_20240329_02",
        "PRED_20240308_01",
        datetime(2024, 3, 29, 18, 0, 0),
        datetime(2024, 3, 8, 0, 0, 0),
        datetime(2024, 3, 29, 0, 0, 0),
        "数据/原始/01_期货行情.csv;数据/原始/06_库存.csv;数据/原始/07_成本利润.csv",
        "螺纹05合约从3691破位大跌至3412（-7.6%），铁矿从877暴跌至740（-15.6%），焦炭跌破2000，负反馈踩踏完美兑现",
        "完全正确",
        "未交易(基准模拟全产业链做空大丰收)",
        "无逻辑错误",
        "无",
        "终端真实表需断崖式走弱，钢厂亏损突破极限引发原料抛售踩踏，成本支撑全面塌陷，产业链完成从隐性背离到共振杀跌的因果闭环",
        "无",
        "将原料负反馈加速度与钢厂现金流亏损深度直接联动纳入主跌浪持仓扩展模块",
        "ARCH-H-005;H-002",
        None,
        "V0.2",
        "已结项"
    ]

    for col_idx, val in enumerate(rev_1, start=1):
        ws.cell(row=6, column=col_idx, value=val)
    for col_idx, val in enumerate(rev_2, start=1):
        ws.cell(row=7, column=col_idx, value=val)

    wb.save(A08_PATH)
    print("    -> A08 成功写入 2 条复盘归因记录 (Row 6, Row 7)")

def update_a09():
    print("[+] 更新 A09_研究假设库.xlsx 历史验证结果 ...")
    wb = openpyxl.load_workbook(A09_PATH)
    ws = wb['假设登记']

    # 查找并更新
    # Row 7: H-002
    # Row 8: H-003
    # Row 13: ARCH-H-005
    # Row 15: ARCH-H-007
    updates = {
        "H-002": {
            "col_h": "2024Q1 (2024-01-01 至 2024-03-31)",
            "col_i": "历史验证通过",
            "col_q": "2024Q1推演表明：被动累库与需求断崖主导的累库带来剧烈下行，而供需双弱下的被动去库无法改变价格颓势，验证了供需结构来源决定后续走势。"
        },
        "H-003": {
            "col_h": "2024Q1 (2024-01-01 至 2024-03-31)",
            "col_i": "历史验证通过",
            "col_q": "2024年1月中下旬盘面维持高估值而基本面实物恶化，背离状态下模型通过降低置信度(Medium)成功规避左侧踩踏，并在3月现实证伪后完成共振捕捉。"
        },
        "ARCH-H-005": {
            "col_h": "2024Q1 (2024-01-01 至 2024-03-31)",
            "col_i": "历史验证通过",
            "col_q": "2024Q1螺纹表需(R04)与库存(R02+R03)分解清晰表明，去库迟滞是由于终端开工断崖而非产量激增所致，实物驱动力拆分有效。"
        },
        "ARCH-H-007": {
            "col_h": "2024Q1 (2024-01-01 至 2024-03-31)",
            "col_i": "历史验证通过",
            "col_q": "2024Q1证实原料与成材关系高度依赖产业状态：节前成材亏损向原料传导存在时滞，3月一旦钢厂减产触发展开负反馈，原料(铁矿跌26%)跌幅远超成材(跌15%)。"
        }
    }

    for r in range(6, 25):
        h_id = ws.cell(row=r, column=1).value
        if h_id in updates:
            u = updates[h_id]
            # column 8: 历史样本范围
            ws.cell(row=r, column=8, value=u["col_h"])
            # column 9: 历史验证结果
            ws.cell(row=r, column=9, value=u["col_i"])
            # column 17: 备注
            ws.cell(row=r, column=17, value=u["col_q"])
            print(f"    -> 更新假设 {h_id}: {u['col_i']}")

    wb.save(A09_PATH)
    print("    -> A09 假设库验证状态更新成功！")

def main():
    update_a07()
    update_a08()
    update_a09()
    print("=" * 80)
    print("三大正式量化账本 (A07, A08, A09) 更新落盘成功！")
    print("=" * 80)

if __name__ == "__main__":
    main()
